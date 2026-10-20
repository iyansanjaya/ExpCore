import csv
import json
import os
import subprocess
import sys
import tempfile
from contextlib import nullcontext
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, Mock, patch

import pdfplumber
from openpyxl import load_workbook

from ExpCore import ExpCore
from tests.fixtures import BPBS_A3_NAME, BPBS_NAME, BPPU_A2_NAME, BPPU_LINES, BPBS_LINES, BPPU_NAME, write_all

ROOT = Path(__file__).resolve().parent


def run_engine(*args):
    result = subprocess.run([sys.executable, str(ROOT / "expcore_engine.py"), *args],
                            capture_output=True, timeout=120)
    # Protokol wajib ASCII murni; decode("ascii") gagal bila ada keluaran liar.
    events = [json.loads(line) for line in result.stdout.decode("ascii").splitlines()]
    return result.returncode, events, result.stderr.decode("utf-8", "replace")


def sheet_row(path, rows=1):
    """Baris data pertama sheet aktif; sheet harus berisi tepat `rows` baris data."""
    book = load_workbook(path)
    try:
        sheet = book.active
        assert sheet.max_row == rows + 1, sheet.max_row
        return dict(zip((c.value for c in sheet[1]), (c.value for c in sheet[2])))
    finally:
        book.close()


# Acuan Excel PDF contoh: lokal saja (contoh-pdf berisi data nasabah asli dan ada di .gitignore).
ACUAN = ROOT / "contoh-pdf" / "acuan"
# Folder contoh -> (metode ExpCore, nama Excel hasil). Pajak Masukan belum punya faktur asli di contoh-pdf;
# tambahkan ("process_pm", "Hasil_Pajak_Masukan.xlsx") bila foldernya ada.
SAMPLE_EXPORTS = {
    "BUPOT 2026": ("process_bupot", "!Hasil_Rekap_Bupot.xlsx"),
    "BUPOT 2024": ("process_bupot_2024", "!Hasil_Rekap_Bupot_2024.xlsx"),
    "Rekening Giro": ("process_rekening_koran", "!Hasil_Rekap_Rekening_Koran.xlsx"),
    "Rekening Tahapan": ("process_rekening_koran", "!Hasil_Rekap_Rekening_Koran.xlsx"),
    "BANG CIMB": ("process_rekening_cimb", "!Hasil_Rekap_Rekening_CIMB.xlsx"),
}


def gaya_sel(cell):
    """Isi dan tampilan satu sel. Nilai dibandingkan dengan ==, jadi 1 dan 1.0 dianggap sama."""
    font, align, border, link = cell.font, cell.alignment, cell.border, cell.hyperlink
    return (cell.value, cell.number_format, font.b, font.i, font.u, font.name, font.sz,
            font.color.rgb if font.color else None, cell.fill.fill_type, cell.fill.fgColor.rgb,
            align.horizontal, align.vertical, align.wrap_text,
            border.left.style, border.right.style, border.top.style, border.bottom.style,
            (link.location, link.target) if link else None)


def bandingkan_excel(acuan, hasil):
    """Maks. 10 perbedaan dua workbook, kosong bila identik: sheet, dimensi, freeze, filter, lebar kolom, setiap sel."""
    a_book, b_book = load_workbook(acuan), load_workbook(hasil)  # File langsung ditutup setelah dibaca.
    if a_book.sheetnames != b_book.sheetnames:
        return [f"sheet {a_book.sheetnames} vs {b_book.sheetnames}"]
    beda = []
    for a, b in zip(a_book.worksheets, b_book.worksheets):
        for label, x, y in (("dimensi", a.dimensions, b.dimensions), ("freeze", a.freeze_panes, b.freeze_panes),
                            ("filter", a.auto_filter.ref, b.auto_filter.ref)):
            if x != y:
                beda.append(f"{a.title} {label}: {x!r} vs {y!r}")
        for kolom in sorted(set(a.column_dimensions) | set(b.column_dimensions)):
            x, y = a.column_dimensions[kolom].width, b.column_dimensions[kolom].width
            if x != y:
                beda.append(f"{a.title} lebar kolom {kolom}: {x} vs {y}")
        for row_a, row_b in zip(a.iter_rows(), b.iter_rows()):
            for x, y in zip(row_a, row_b):
                if gaya_sel(x) != gaya_sel(y):
                    beda.append(f"{a.title}!{x.coordinate}: {gaya_sel(x)} vs {gaya_sel(y)}")
    return beda[:10]


def test_engine_protocol():
    """Engine JSON: lima ekspor + penamaan pada PDF sintetis, folder Unicode, PDF rusak, argumen."""
    with tempfile.TemporaryDirectory() as temp:
        folder = Path(temp) / "Bupot Ünïcode — 2026 (uji)"
        write_all(folder / "sub")
        (folder / "rusak.pdf").write_bytes(b"corrupt")
        single = "Selesai — 1 baris, 5 PDF dilewati."
        exports = {
            "bupot": ("!Hasil_Rekap_Bupot.xlsx", single, {"Nomor Dokumen": "25004WOBY", "DPP (Rp)": 3700000}),
            "bupot2024": ("!Hasil_Rekap_Bupot_2024.xlsx", single, {"Nomor Bukti Potong": "2000000001", "Tarif (%)": 2}),
            "pm": ("Hasil_Pajak_Masukan.xlsx", single,
                   {"NPWP Pembeli": "0012345678901234", "DPP": 3000000, "PPN": 360000}),
            "rekening": ("!Hasil_Rekap_Rekening_Koran.xlsx",
                         "Selesai — 1 rekening koran, 5 transaksi, 0 perlu dicek, 5 PDF dilewati.",
                         {"No. Rekening": "1234567890", "Mutasi CR": 3736425.01, "Hasil Cek": "SESUAI"}),
            # Ringkasan CIMB: satu baris per rekening (4 rekening dalam satu laporan).
            "cimb": ("!Hasil_Rekap_Rekening_CIMB.xlsx",
                     "Selesai — 1 laporan, 4 rekening, 5 transaksi, 0 perlu dicek, 5 PDF dilewati.",
                     {"No. Rekening": "1111222233", "Saldo Akhir": 1234.5, "Hasil Cek": "SESUAI"}, 4),
        }
        for job, (output, summary, expected, *rows) in exports.items():
            code, events, stderr = run_engine(job, str(folder))
            assert code == 0, (job, events, stderr)
            assert events[0] == {"event": "log", "message": "Memproses 6 file …"}, events[0]
            progress = [e for e in events if e["event"] == "progress"]
            assert progress == [{"event": "progress", "done": i, "total": 6} for i in range(6)], progress
            assert sum(e["event"] in ("done", "error") for e in events) == 1, events
            assert events[-1]["event"] == "done" and events[-1]["path"] == str(folder / output), events[-1]
            assert events[-1]["summary"] == summary, events[-1]
            values = sheet_row(events[-1]["path"], *rows)
            assert values["Folder Sumber"] == "sub", values
            for column, value in expected.items():
                assert values[column] == value, (job, column, values[column])

        code, events, _ = run_engine("rename", str(folder))
        # BPPU (2026) dan BPBS (2024) dalam satu folder; faktur dan rekening koran BCA/CIMB dilewati.
        assert code == 0 and events[-1]["summary"] == (
            "Pratinjau selesai — 2 siap/berhasil, 0 sudah sesuai, 3 perlu diperiksa, 1 gagal."), events[-1]
        assert events[0]["message"] == "Pratinjau: memeriksa 6 PDF … (nama dari Identitas Pemotong)", events[0]
        assert any(e.get("message", "").endswith(f"{BPBS_NAME} [BPBS]") for e in events), events
        assert (folder / "sub" / "bppu.pdf").exists() and (folder / "sub" / "bpbs.pdf").exists(), \
            "Pratinjau tidak boleh mengubah nama file"
        code, events, _ = run_engine("rename", str(folder), "--apply")
        assert code == 0 and events[-1]["summary"].startswith("Penerapan selesai — 2 siap/berhasil"), events[-1]
        for old, new in (("bppu.pdf", BPPU_NAME), ("bpbs.pdf", BPBS_NAME)):
            assert (folder / "sub" / new).exists() and not (folder / "sub" / old).exists(), new
        assert Path(events[-1]["path"]).name.startswith("Log_Penamaan_Bupot_Penerapan_")
        code, events, _ = run_engine("rename", str(folder))
        assert "0 siap/berhasil, 2 sudah sesuai" in events[-1]["summary"], events[-1]
        # Sumber nama wajib pajak yang dipotong: A.2 pada BPPU, A.3 pada BPBS; file tidak berubah.
        code, events, _ = run_engine("rename", str(folder), "--nama", "penerima")
        assert code == 0 and "2 siap/berhasil, 0 sudah sesuai" in events[-1]["summary"], events[-1]
        with open(events[-1]["path"], encoding="utf-8-sig", newline="") as log:
            ready = [row for row in csv.DictReader(log) if row["status"] == "SIAP"]
        assert [(row["nama_baru"], row["sumber_nama"], row["FORMULIR"]) for row in ready] == [
            (BPPU_A2_NAME, "Wajib Pajak yang Dipotong (A.2)", "BPPU"),
            (BPBS_A3_NAME, "Wajib Pajak yang Dipotong (A.3)", "BPBS")], ready
        assert (folder / "sub" / BPPU_NAME).exists() and (folder / "sub" / BPBS_NAME).exists()

        # Kegagalan menulis hasil menjadi event error, bukan crash diam-diam.
        (folder / "!Hasil_Rekap_Bupot.xlsx").unlink()
        (folder / "!Hasil_Rekap_Bupot.xlsx").mkdir()
        code, events, stderr = run_engine("bupot", str(folder))
        assert code == 1 and events[-1]["event"] == "error" and "Traceback" in stderr, (events, stderr)
        assert events[-2]["message"].startswith("Error: "), events[-2]

        empty = Path(temp) / "kosong"
        empty.mkdir()
        for path, message in ((empty, "Tidak ada PDF di folder atau subfolder ini. Pilih folder lain."),
                              (Path(temp) / "tidak-ada", "Folder tidak ditemukan. Pilih folder yang tersedia.")):
            assert run_engine("pm", str(path))[:2] == (1, [{"event": "error", "message": message}])
        for args in (("bupot", str(folder), "--apply"), ("lainnya", str(folder)), ("bupot",),
                     ("rekening", str(folder), "--nama", "penerima"), ("cimb", str(folder), "--apply"),
                     ("rename", str(folder), "--nama", "C.3")):
            code, events, stderr = run_engine(*args)
            assert (code, events) == (2, []) and "usage" in stderr, (args, code, events)

    # Proses pekerja pembaca PDF mewarisi stdout engine di Windows; keluarannya wajib ke stderr.
    code = "from ExpCore import ExpCore\nlist(ExpCore(workers=2)._peta_pdf(print, ['pekerja-a', 'pekerja-b']))"
    result = subprocess.run([sys.executable, "-c", code], cwd=ROOT, capture_output=True, timeout=120)
    assert result.returncode == 0 and result.stdout == b"", (result.stdout, result.stderr)
    assert b"pekerja-a" in result.stderr and b"pekerja-b" in result.stderr, result.stderr
    print("Engine: protokol JSON, lima ekspor, penamaan BPPU + BPBS (kedua sumber nama), Unicode, PDF rusak, "
          "argumen, dan stdout proses pekerja OK")


def test_batch_exports():
    """All three workers export valid PDFs despite a corrupt sibling PDF."""
    for method, parser in (
        ("process_bupot", "_extract_bupot_rows"),
        ("process_bupot_2024", "_extract_bupot2024_rows"),
        ("process_pm", None),
    ):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder)
            (source / "sub").mkdir()
            (source / "broken.pdf").write_bytes(b"corrupt")
            (source / "sub" / "valid.pdf").write_bytes(b"fixture")
            page = SimpleNamespace(
                extract_text=lambda: "Pembeli Barang Kena Pajak\nNama : PT UJI\nNPWP : 0012345678901234\n"
                                     "Kode dan Nomor Seri Faktur Pajak : 0100000000000001",
                extract_tables=lambda: [[["1", "123456", "Jasa Uji Rp 1.000,00 x 2 Unit"]]],
            )
            # Satu proses (workers=1): patch tidak ikut ke proses pekerja.
            worker = ExpCore(log=Mock(), progress=Mock(), workers=1)
            parsed = [{"NPWP": "0012345678901234", "DPP (Rp)": 2000.0}]
            good_pdf = MagicMock()
            good_pdf.__enter__.return_value = SimpleNamespace(pages=[page])
            parser_patch = patch.object(ExpCore, parser, staticmethod(lambda text: parsed)) if parser else nullcontext()
            with patch("ExpCore.pdfplumber.open", side_effect=[ValueError("corrupt PDF"), good_pdf]), parser_patch:
                output, summary = getattr(worker, method)(folder)
            assert "1 baris, 1 PDF dilewati" in summary, summary
            book = load_workbook(output)
            sheet = book.active
            assert sheet.max_row == 2
            values = dict(zip((c.value for c in sheet[1]), (c.value for c in sheet[2])))
            assert values["Folder Sumber"] == "sub"
            if method == "process_pm":
                assert values["NPWP Pembeli"] == "0012345678901234"
                assert values["DPP"] == 2000 and values["PPN"] == 240
            else:
                assert values["NPWP"] == "0012345678901234"
            book.close()
            assert (source / "broken.pdf").read_bytes() == b"corrupt"
            assert (source / "sub" / "valid.pdf").read_bytes() == b"fixture"
            assert worker.progress.call_count == 2
    print("Batch workers: corrupted PDF recovery, subfolders and three Excel exports OK")


def test_rename_pemotong():
    # Teks berdasarkan contoh BPPU: nama A.2 berbeda dari pemotong C.3.
    text = """26007GORO 01-2026 FINAL NORMAL
A.1 NPWP / NIK : 0026264515077000
A.2 NAMA : PLAZA LIFESTYLE PRIMA
A.3 NOMOR IDENTITAS TEMPAT KEGIATAN USAHA (NITKU) : 0026264515077000000000 - PLAZA LIFESTYLE PRIMA
C.3 NAMA PEMOTONG DAN/ATAU PEMUNGUT : VOLANS
PPh
C.4 TANGGAL : 31 Januari 2026"""
    data = ExpCore._extract_rename_bupot_data(text)
    assert data["NAMA_PENERIMA"] == "PLAZA LIFESTYLE PRIMA", data
    assert data["NAMA_PEMOTONG"] == "VOLANS", data

    wrapped = text.replace("A.2 NAMA : PLAZA LIFESTYLE PRIMA",
                           "A.2 NAMA :\nPLAZA LIFESTYLE\nPRIMA")
    assert ExpCore._extract_rename_bupot_data(wrapped)["NAMA_PENERIMA"] == \
        "PLAZA LIFESTYLE PRIMA"
    blank = text.replace("A.2 NAMA : PLAZA LIFESTYLE PRIMA", "A.2 NAMA :")
    absent = text.replace("A.2 NAMA : PLAZA LIFESTYLE PRIMA\n", "")
    for incomplete in (blank, absent):
        assert ExpCore._extract_rename_bupot_data(incomplete)["NAMA_PENERIMA"] == ""

    without_pemotong = text[:text.index("C.3")]
    blank_pemotong = text.replace(": VOLANS", ":")
    wrapped_pemotong = text.replace(": VOLANS", ":\nVOLANS")
    # Sumber nama -> (field, label CSV, nama file, teks yang nama sumbernya kosong). Tanpa fallback:
    # A.2 kosong tidak boleh diganti C.3, dan sebaliknya.
    sources = {
        "pemotong": ("NAMA_PEMOTONG", "Identitas Pemotong", "VOLANS - 26007GORO - 01-2026 - FINAL - NORMAL.pdf",
                     (without_pemotong, blank_pemotong)),
        "penerima": ("NAMA_PENERIMA", "Wajib Pajak yang Dipotong",
                     "PLAZA LIFESTYLE PRIMA - 26007GORO - 01-2026 - FINAL - NORMAL.pdf", (blank, absent)),
    }
    for source, (field, label, expected_name, incomplete) in sources.items():
        for pdf_text in (text, wrapped, wrapped_pemotong, without_pemotong, blank_pemotong, blank, absent):
            complete = pdf_text not in incomplete
            for apply_changes in (False, True):
                with tempfile.TemporaryDirectory() as folder:
                    original = os.path.join(folder, "original.pdf")
                    with open(original, "wb") as pdf_file:
                        pdf_file.write(b"test PDF placeholder")
                    app = SimpleNamespace(
                        progress=Mock(), log=Mock(),
                        _extract_rename_bupot_data=ExpCore._extract_rename_bupot_data,
                        _safe_filename=ExpCore._safe_filename,
                        _unique_filename=ExpCore._unique_filename,
                    )
                    pdf = SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda: pdf_text)])
                    with patch("ExpCore.pdfplumber.open") as open_pdf:
                        open_pdf.return_value.__enter__.return_value = pdf
                        result_path, summary = ExpCore.process_rename_bupot(
                            app, folder, apply_changes=apply_changes, name_source=source)
                        assert Path(result_path).is_file()
                        assert "selesai" in summary
                    assert label in app.log.call_args_list[0].args[0], app.log.call_args_list[0]

                    csv_name, = [name for name in os.listdir(folder) if name.endswith(".csv")]
                    with open(os.path.join(folder, csv_name), encoding="utf-8-sig", newline="") as log:
                        row, = list(csv.DictReader(log))
                    assert row["sumber_nama"] == f"{label} ({'C.3' if source == 'pemotong' else 'A.2'})", row
                    assert row["FORMULIR"] == "", row
                    assert row["NAMA_PEMOTONG"] == ("" if pdf_text in (without_pemotong, blank_pemotong) else "VOLANS")
                    assert row["NAMA_PENERIMA"] == ("" if pdf_text in (blank, absent) else "PLAZA LIFESTYLE PRIMA"), row
                    if complete:
                        assert row["nama_baru"] == expected_name, (source, row)
                        assert row["status"] == ("BERHASIL" if apply_changes else "SIAP"), row
                        assert row["data_tidak_lengkap"] == "", row
                    else:
                        assert row["status"] == ("DILEWATI" if apply_changes else "PERLU CEK"), row
                        assert row["data_tidak_lengkap"] == field, row
                    assert os.path.exists(original) == (not (complete and apply_changes))
                    assert os.path.exists(os.path.join(folder, expected_name)) == (complete and apply_changes)

    try:
        ExpCore.process_rename_bupot(SimpleNamespace(), ".", name_source="C.3")
        raise AssertionError("sumber nama tidak dikenal harus ditolak")
    except ValueError as error:
        assert "Sumber nama tidak dikenal" in str(error)

    print("Rename Bupot: sumber nama C.3 dan A.2 (pratinjau dan penerapan, tanpa fallback): ok")


def test_rename_bpbs():
    """Penamaan Bupot membaca BPBS (Bupot 2024) di samping BPPU; data BPBS yang ambigu tidak ditebak."""
    bpbs = "\n".join(BPBS_LINES)
    base = {"NAMA_PENERIMA": "MITRACOLL SARANA JAYA", "NAMA_PEMOTONG": "BFI FINANCE INDONESIA TBK.",
            "NOMOR_BUKTI": "2000000001", "MASA_PAJAK": "01-2024", "SIFAT": "TIDAK FINAL", "STATUS": "NORMAL",
            "FORMULIR": "BPBS"}
    assert ExpCore._extract_rename_bupot_data(bpbs) == base, ExpCore._extract_rename_bupot_data(bpbs)
    assert ExpCore._extract_rename_bupot_data("\n".join(BPPU_LINES))["FORMULIR"] == "BPPU"

    row = "1-2024 24-104-18 1.300.000,00 2.00 26.000,00"
    for name, old, new, changed in (
        # Status hanya dari kotak H.2/H.3 yang terbaca.
        ("pembetulan", "Pembetulan Ke- 0", "Pembetulan Ke- 2", {"STATUS": "PEMBETULAN KE-2"}),
        ("pembatalan", "H.3 Pembatalan", "H.3 X Pembatalan", {"STATUS": "DIBATALKAN"}),
        ("H.2 tak terbaca", "Pembetulan Ke- 0", "Pembetulan Ke-", {"STATUS": ""}),
        # Sifat: tepat satu kotak H.4/H.5 harus tercentang.
        ("final", "H.4 PPh Final", "H.4 X PPh Final", {"SIFAT": ""}),
        ("tanpa centang", "H.5 X PPh Tidak Final", "H.5 PPh Tidak Final", {"SIFAT": ""}),
        # Masa: satu masa untuk semua baris; masa berbeda tidak dipilih salah satunya.
        ("dua baris satu masa", row, f"{row}\n1-2024 24-104-18 200.000,00 2 4.000,00", {}),
        ("dua masa", row, f"{row}\n2-2024 24-104-18 200.000,00 2 4.000,00", {"MASA_PAJAK": ""}),
        ("tanpa baris", row, "", {"MASA_PAJAK": ""}),
        ("A.3 kosong", "A.3 Nama : MITRACOLL SARANA JAYA", "A.3 Nama :", {"NAMA_PENERIMA": ""}),
        ("C.2 kosong", ": BFI FINANCE INDONESIA TBK.", ":", {"NAMA_PEMOTONG": ""}),
    ):
        variant = bpbs.replace(old, new)
        assert variant != bpbs, name
        assert ExpCore._extract_rename_bupot_data(variant) == {**base, **changed}, (name, ExpCore._extract_rename_bupot_data(variant))
    # Hanya H.4 yang tercentang -> FINAL.
    final_only = bpbs.replace("H.4 PPh Final", "H.4 X PPh Final").replace("H.5 X PPh Tidak Final", "H.5 PPh Tidak Final")
    assert ExpCore._extract_rename_bupot_data(final_only)["SIFAT"] == "FINAL"
    # Rekap Bupot 2024 memakai pembaca BPBS yang sama; rekap Bupot 2026 tetap tidak membaca BPBS.
    assert len(ExpCore._extract_bupot2024_rows(bpbs)) == 1 and ExpCore._extract_bupot_rows(bpbs) == []

    with tempfile.TemporaryDirectory() as temp:
        folder = Path(temp)
        write_all(folder)
        write_all(folder / "salinan")  # BPBS kedua dengan data sama -> nama unik "(2)" di foldernya sendiri
        for source, expected in (("pemotong", {"bppu.pdf": BPPU_NAME, "bpbs.pdf": BPBS_NAME}),
                                 ("penerima", {"bppu.pdf": BPPU_A2_NAME, "bpbs.pdf": BPBS_A3_NAME})):
            log_path, summary = ExpCore().process_rename_bupot(str(folder), name_source=source)
            assert summary == "Pratinjau selesai — 4 siap/berhasil, 0 sudah sesuai, 6 perlu diperiksa, 0 gagal.", summary
            with open(log_path, encoding="utf-8-sig", newline="") as log:
                rows = {(r["folder_sumber"], r["nama_lama"]): r for r in csv.DictReader(log)}
            codes = {"pemotong": {"BPPU": "C.3", "BPBS": "C.2"}, "penerima": {"BPPU": "A.2", "BPBS": "A.3"}}[source]
            for sub in (".", "salinan"):
                for old, new in expected.items():
                    r = rows[(sub, old)]
                    form = "BPBS" if old == "bpbs.pdf" else "BPPU"
                    assert (r["status"], r["nama_baru"], r["FORMULIR"]) == ("SIAP", new, form), r
                    assert r["sumber_nama"].endswith(f"({codes[form]})"), r
                for old in ("faktur.pdf", "rekening.pdf", "cimb.pdf"):
                    assert rows[(sub, old)]["status"] == "PERLU CEK" and rows[(sub, old)]["FORMULIR"] == "", rows[(sub, old)]
        _, summary = ExpCore().process_rename_bupot(str(folder), apply_changes=True)
        assert summary.startswith("Penerapan selesai — 4 siap/berhasil"), summary
        for sub in (folder, folder / "salinan"):
            assert (sub / BPPU_NAME).exists() and (sub / BPBS_NAME).exists(), sorted(p.name for p in sub.iterdir())
            assert not (sub / "bppu.pdf").exists() and not (sub / "bpbs.pdf").exists()
        # Nama hasil sudah sesuai: penerapan ulang tidak mengubah apa pun.
        _, summary = ExpCore().process_rename_bupot(str(folder), apply_changes=True)
        assert summary.startswith("Penerapan selesai — 0 siap/berhasil, 4 sudah sesuai"), summary
        # Dua BPBS identik di folder yang sama: yang kedua mendapat akhiran (2), tidak saling menimpa.
        dup = folder / "ganda"
        write_all(dup)
        (dup / "bpbs.pdf").replace(dup / "a.pdf")
        write_all(dup / "tmp")
        (dup / "tmp" / "bpbs.pdf").replace(dup / "b.pdf")
        ExpCore().process_rename_bupot(str(dup), apply_changes=True)
        assert (dup / BPBS_NAME).exists() and (dup / BPBS_NAME.replace(".pdf", " (2).pdf")).exists()
    print("Penamaan Bupot BPBS (Bupot 2024): deteksi formulir, status/sifat/masa tanpa tebakan, nama unik: ok")


def read_rekening(path):
    with pdfplumber.open(path) as pdf:
        return ExpCore._extract_rekening_koran(pdf.pages)


def test_rekening_koran():
    """e-Statement BCA: kolom dari posisi x, debit bertanda DB, ringkasan terpisah dan dicocokkan."""
    from datetime import datetime
    from decimal import Decimal

    from tests.fixtures import REKENING_LINES, REKENING_SUMMARY, write_rekening

    with tempfile.TemporaryDirectory() as temp:
        temp = Path(temp)

        def variant(name, **options):
            path = temp / f"{name}.pdf"
            write_rekening(path, **options)
            return read_rekening(path)

        data = variant("asli")
        assert {k: data[k] for k in ("jenis", "norek", "nama", "periode", "mata_uang", "catatan")} == {
            "jenis": "REKENING GIRO", "norek": "1234567890", "nama": "PT CONTOH SEJAHTERA",
            "periode": "JANUARI 2025", "mata_uang": "IDR", "catatan": []}, data
        D = Decimal
        assert data["transaksi"] == [
            {"tanggal": datetime(2025, 1, 2), "keterangan": "SETORAN TUNAI", "cbg": "0473",
             "debit": None, "kredit": D("2500000.00"), "saldo": None},
            {"tanggal": datetime(2025, 1, 2), "keterangan": "TRSF E-BANKING DB 0201/FTSCY/WS95051 750000.00 BUDI SANTOSO",
             "cbg": "", "debit": D("750000.00"), "kredit": None, "saldo": D("2750000.00")},
            # Keterangan berlanjut ke halaman 2 tetap milik transaksi yang sama.
            {"tanggal": datetime(2025, 1, 3), "keterangan": "TRSF E-BANKING CR 0301/FTSCY/WS95031 1234567.89 PT MITRA UJI",
             "cbg": "", "debit": None, "kredit": D("1234567.89"), "saldo": D("3984567.89")},
            {"tanggal": datetime(2025, 1, 31), "keterangan": "BUNGA", "cbg": "",
             "debit": None, "kredit": D("1857.12"), "saldo": None},
            {"tanggal": datetime(2025, 1, 31), "keterangan": "PAJAK BUNGA", "cbg": "",
             "debit": D("371.42"), "kredit": None, "saldo": D("3986053.59")},
        ], data["transaksi"]
        assert data["ringkasan"] == {"SALDO AWAL": (D("1000000.00"), None), "MUTASI CR": (D("3736425.01"), 3),
                                     "MUTASI DB": (D("750371.42"), 2), "SALDO AKHIR": (D("3986053.59"), None)}

        # Mutasi bulan sebelumnya pada periode Januari milik tahun sebelumnya.
        rollover = [line if line is None or line[0] != "02/01" or line[1] != "SETORAN TUNAI"
                    else ("31/12", *line[1:]) for line in REKENING_LINES]
        data = variant("tahun", lines=rollover)
        assert data["transaksi"][0]["tanggal"] == datetime(2024, 12, 31) and data["catatan"] == [], data

        # Setiap ketidakcocokan tercatat; tidak ada yang lolos sebagai SESUAI.
        no_db = [line if line is None or line[1] != "PAJAK BUNGA" else (*line[:5], "", line[6])
                 for line in REKENING_LINES]
        wrong_total = [row if row[0] != "MUTASI CR" else ("MUTASI CR", "3,736,425.00", "3") for row in REKENING_SUMMARY]
        odd_saldo = [line if line is None or line[1] != "PAJAK BUNGA" else (*line[:6], "-3,986,053.59")
                     for line in REKENING_LINES]
        for name, options, expected in (
            ("tanpa-db", {"lines": no_db}, ["Kredit terbaca 3,736,796.43 dari 4 transaksi, ringkasan PDF "
                                            "3,736,425.01 dari 3.", "Debit terbaca 750,000.00 dari 1 transaksi, "
                                            "ringkasan PDF 750,371.42 dari 2.",
                                            "31/01: saldo PDF 3,986,053.59, hasil hitung 3,986,796.43."]),
            ("total", {"summary": wrong_total}, ["Kredit terbaca 3,736,425.01 dari 3 transaksi, ringkasan PDF "
                                                 "3,736,425.00 dari 3.", "Ringkasan PDF tidak seimbang: saldo awal "
                                                 "+ mutasi CR − mutasi DB ≠ saldo akhir."]),
            ("tanpa-ringkasan", {"summary": []}, ["Ringkasan saldo dan mutasi di akhir PDF tidak lengkap."]),
            ("halaman", {"total_pages": 3}, ["PDF hanya berisi 2 dari 3 halaman."]),
            ("saldo-aneh", {"lines": odd_saldo}, ["Halaman 2: teks tak dikenal di kolom tanggal/angka: -3,986,053.59"]),
        ):
            assert variant(name, **options)["catatan"] == expected, (name, variant(name, **options)["catatan"])

        # PDF lain (bukan e-Statement BCA) dilewati, termasuk laporan CIMB Niaga.
        write_all(temp / "lain")
        for name in ("bppu", "bpbs", "faktur", "cimb"):
            assert read_rekening(temp / "lain" / f"{name}.pdf") is None, name

        # Rekap folder: Ringkasan + sheet per PDF urut rekening/periode, PDF lain & rusak dilewati.
        folder = temp / "Rekening Ünïcode"
        write_rekening(folder / "b" / "feb.pdf", periode="FEBRUARI 2025")
        write_rekening(folder / "a" / "jan.pdf")
        write_rekening(folder / "a" / "jan salinan.pdf")
        write_rekening(folder / "tahapan.pdf", title="REKENING TAHAPAN", summary=wrong_total)
        write_all(folder / "lain")
        (folder / "lain" / "rekening.pdf").unlink()
        (folder / "rusak.pdf").write_bytes(b"corrupt")
        app = ExpCore(log=Mock(), progress=Mock())
        output, summary = app.process_rekening_koran(str(folder))
        # Periode yang sama di dua PDF: keduanya ditandai PERLU CEK (tidak ada yang dibuang).
        assert summary == "Selesai — 4 rekening koran, 20 transaksi, 3 perlu dicek, 5 PDF dilewati.", summary
        logs = [call.args[0] for call in app.log.call_args_list]
        assert ("PERINGATAN: REKENING GIRO 1234567890 periode 01-2025 muncul di 2 PDF: jan salinan.pdf, jan.pdf"
                " — ditandai PERLU CEK." in logs), logs
        assert "DILEWATI: lain\\faktur.pdf — bukan e-Statement Rekening Giro/Tahapan BCA." in logs, logs

        book = load_workbook(output)
        try:
            assert book.sheetnames == ["Ringkasan", "Giro 1234567890 2025-01", "Giro 1234567890 2025-01 (2)",
                                       "Giro 1234567890 2025-02", "Tahapan 1234567890 2025-01"], book.sheetnames
            overview = [[cell.value for cell in row] for row in book["Ringkasan"].iter_rows()]
            assert overview[0][:8] == ["No", "Sheet", "Jenis Rekening", "No. Rekening", "Nama", "Periode",
                                       "Mata Uang", "Saldo Awal"], overview[0]
            assert overview[1] == [1, "Giro 1234567890 2025-01", "REKENING GIRO", "1234567890", "PT CONTOH SEJAHTERA",
                                   "JANUARI 2025", "IDR", 1000000, 3736425.01, 3, 750371.42, 2, 3986053.59, 5,
                                   "PERLU CEK", "Laporan periode yang sama juga ada di a\\jan.pdf; transaksinya bisa tercatat ganda.", "a", "jan salinan.pdf"], overview[1]
            assert [row[14] for row in overview[1:]] == ["PERLU CEK", "PERLU CEK", "SESUAI", "PERLU CEK"]
            assert overview[4][15].startswith("Kredit terbaca 3,736,425.01 dari 3 transaksi"), overview[4]
            assert book["Ringkasan"]["B2"].hyperlink.location == "'Giro 1234567890 2025-01'!A1"

            sheet = book["Giro 1234567890 2025-02"]
            table = [[cell.value for cell in row[:6]] for row in sheet.iter_rows(max_row=6)]
            assert table[0] == ["TANGGAL", "KETERANGAN", "CBG", "DEBIT", "CREDIT", "SALDO"]
            assert table[1] == [datetime(2025, 1, 2), "SETORAN TUNAI", "0473", None, 2500000, None], table[1]
            assert table[5] == [datetime(2025, 1, 31), "PAJAK BUNGA", None, 371.42, None, 3986053.59], table[5]
            assert sheet.max_row == 16 and sheet["A7"].value is None, "ringkasan tidak boleh di bawah tabel"
            assert (sheet["A2"].number_format, sheet["C2"].number_format, sheet["D3"].number_format) == (
                "dd/mm/yyyy", "@", "#,##0.00")
            assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:F6"
            panel = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=9, max_row=16, min_col=8, max_col=10)]
            assert panel == [["RINGKASAN PDF", "NILAI", "TRANSAKSI"], ["SALDO AWAL", 1000000, None],
                             ["MUTASI CR", 3736425.01, 3], ["MUTASI DB", 750371.42, 2],
                             ["SALDO AKHIR", 3986053.59, None], [None, None, None],
                             ["HASIL CEK", None, None], ["SESUAI", None, None]], panel
            assert sheet["I3"].value == "1234567890" and sheet["I5"].value == "FEBRUARI 2025"
            # _tulis_panel memformat nomor rekening sebagai teks dari label persisnya (dipakai juga oleh CIMB).
            assert sheet["I3"].number_format == "@" and sheet["I10"].number_format == "#,##0.00"
            flagged = book["Tahapan 1234567890 2025-01"]
            assert flagged["H16"].value == "PERLU CEK" and flagged["H17"].value.startswith("• Kredit terbaca")
        finally:
            book.close()

        # Nama sheet maks. 31 karakter: jenis rekening yang dipotong, nomor dan periode tetap utuh.
        long_title = temp / "judul panjang"
        write_rekening(long_title / "x.pdf", title="REKENING TAHAPAN XPRESI PLUS")
        book = load_workbook(ExpCore(log=Mock(), progress=Mock()).process_rekening_koran(str(long_title))[0])
        assert book.sheetnames == ["Ringkasan", "Tahapan Xpre 1234567890 2025-01"], book.sheetnames
        book.close()

        empty = temp / "bukan-rekening"
        write_all(empty)
        (empty / "rekening.pdf").unlink()
        assert ExpCore(log=Mock(), progress=Mock()).process_rekening_koran(str(empty))[0] is None
    print("Rekening koran: kolom posisi, debit DB, lintas halaman, cek ringkasan & saldo, Excel: ok")


def read_cimb(path):
    with pdfplumber.open(path) as pdf:
        return ExpCore._extract_rekening_cimb(pdf.pages)


def test_rekening_cimb():
    """CIMB Niaga: sel multi-baris per garis tabel, kolom dari garis di bawah judul, bagian dari judulnya."""
    from datetime import datetime
    from decimal import Decimal as D

    from tests.fixtures import CIMB_DAFTAR_GIRO, CIMB_JUDUL_TABEL, CIMB_MUTASI, write_cimb

    def ganti(rows, tanggal, line):  # baris bertanggal `tanggal` diganti `line`
        return [row if row is None else [line if part[0] == tanggal else part for part in row] for row in rows]

    with tempfile.TemporaryDirectory() as temp:
        temp = Path(temp)

        def variant(name, **options):
            path = temp / f"{name}.pdf"
            write_cimb(path, **options)
            return read_cimb(path)

        data = variant("asli")
        assert {k: data[k] for k in ("nasabah", "cif", "tanggal_laporan", "periode", "bulan", "tahun")} == {
            "nasabah": "BUDI CONTOH", "cif": "11223344556677", "tanggal_laporan": datetime(2025, 1, 31),
            "periode": "JANUARI 2025", "bulan": 1, "tahun": 2025}, data
        # Urutan daftar RINCIAN ASET; rekening deposito 9999000011 dan Poin Xtra tidak ikut.
        assert [(a["kategori"], a["norek"], a["mata_uang"], a["produk"], a["SALDO AWAL"], a["total"],
                 a["SALDO AKHIR"], a["saldo_daftar"], a["saldo_idr"], len(a["transaksi"]), a["catatan"])
                for a in data["rekening"]] == [
            ("Tabungan/Giro", "1111222233", "IDR", "TAPLUS CONTOH", D("1000.00"), (D("265.50"), D("500.00")),
             D("1234.50"), D("1234.50"), D("1234.50"), 2, []),
            ("Tabungan/Giro", "4444555566", "USD", "VALAS CONTOH USD", D("10.00"), (D(0), D(0)),
             D("10.00"), D("10.00"), D("160000.00"), 0, []),
            ("Tabungan Berjangka", "7777888899", "USD", "RENCANA CONTOH USD", D("100.00"), (D("0.25"), D("5.25")),
             D("105.00"), D("105.00"), D("1680000.00"), 2, []),
            # Jatuh tempo: saldo "-" di daftar aset = 0.
            ("Tabungan Berjangka", "7777000011", "USD", "RENCANA CONTOH USD", D("50.00"), (D("50.00"), D(0)),
             D("0.00"), D(0), D(0), 1, []),
        ], data["rekening"]
        # Uraian beberapa baris (tanggal & nominal di tengah sel) tetap satu transaksi; 31/12 di laporan
        # Januari milik tahun sebelumnya; tabel berlanjut ke halaman berikutnya.
        assert data["rekening"][0]["transaksi"] == [
            {"tanggal": datetime(2024, 12, 31), "valuta": datetime(2024, 12, 31),
             "keterangan": "SETORAN TUNAI CABANG CONTOH REF0001", "cek": "", "debit": None, "kredit": D("500.00"),
             "saldo": D("1500.00")},
            {"tanggal": datetime(2025, 1, 5), "valuta": datetime(2025, 1, 6), "keterangan": "TARIK TUNAI",
             "cek": "CK001", "debit": D("265.50"), "kredit": None, "saldo": D("1234.50")},
        ], data["rekening"][0]["transaksi"]
        assert data["kategori"] == [
            {"kategori": "Tabungan/Giro", "rekening": 2, "total": D("161234.50"), "jumlah": D("161234.50"),
             "catatan": []},
            {"kategori": "Tabungan Berjangka", "rekening": 2, "total": D("1680000.00"), "jumlah": D("1680000.00"),
             "catatan": []},
        ], data["kategori"]

        # Valuta melewati pergantian tahun mengikuti tanggal transaksinya.
        lewat_tahun = ganti(CIMB_MUTASI, "31/12", ("31/12", "02/01", "CABANG CONTOH", "", "", "500.00", "1,500.00"))
        assert variant("valuta", mutasi=lewat_tahun)["rekening"][0]["transaksi"][0]["valuta"] == datetime(2025, 1, 2)

        # Setiap ketidakcocokan tercatat pada rekening/kategori yang terkena; tidak ada yang lolos sebagai SESUAI.
        tertukar = ganti(CIMB_MUTASI, "05/01", ("05/01", "06/01", "TARIK TUNAI", "CK001", "", "265.50", "1,234.50"))
        ganda = ganti(CIMB_MUTASI, "05/01", ("05/01", "06/01", "TARIK TUNAI", "CK001", "265.50", "1.00", "1,234.50"))
        salah_tanggal = ganti(CIMB_MUTASI, "05/01", ("32/01", "06/01", "TARIK TUNAI", "", "265.50", "", "1,234.50"))
        daftar_beda = [("1111222233", "TAPLUS CONTOH", "IDR", "1,234.00", "1,234.00"), CIMB_DAFTAR_GIRO[1]]
        daftar_lebih = [*CIMB_DAFTAR_GIRO, ("5555666677", "TAPLUS LAIN", "IDR", "0.00", "0.00")]
        for name, options, expected, kategori in (
            ("tertukar", {"mutasi": tertukar}, [
                "Debit terbaca 0.00 dari 0 transaksi, baris Total PDF 265.50.",
                "Kredit terbaca 765.50 dari 2 transaksi, baris Total PDF 500.00.",
                "05/01: saldo PDF 1,234.50, hasil hitung 1,765.50."], []),
            ("ganda", {"mutasi": ganda}, [
                "Hal. 3 05/01: susunan kolom tidak dikenal — TARIK TUNAI",
                "Kredit terbaca 501.00 dari 2 transaksi, baris Total PDF 500.00.",
                "05/01: saldo PDF 1,234.50, hasil hitung 1,235.50."], []),
            ("tanggal", {"mutasi": salah_tanggal}, ["Hal. 3 32/01: tanggal transaksi tidak valid (32/01)."], []),
            ("tanpa-saldo-akhir", {"mutasi": CIMB_MUTASI[:-1]}, ["Baris SALDO AKHIR tidak ditemukan."], []),
            ("tanpa-total", {"mutasi": [row for row in CIMB_MUTASI if row is None or row[0][0] != "Total"]},
             ["Baris Total debit/kredit tidak ditemukan."], []),
            ("daftar", {"daftar_giro": daftar_beda},
             ["Saldo akhir 1,234.50 berbeda dari saldo di daftar Tabungan/Giro 1,234.00."],
             ["Jumlah saldo Ekuivalen IDR 161,234.00 berbeda dari total PDF 161,234.50."]),
            ("total-giro", {"total_giro": "161,234.00"}, [],
             ["Jumlah saldo Ekuivalen IDR 161,234.50 berbeda dari total PDF 161,234.00."]),
            ("halaman", {"total_pages": 6}, ["PDF hanya berisi 5 dari 6 halaman."],
             ["PDF hanya berisi 5 dari 6 halaman."]),
        ):
            data = variant(name, **options)
            assert data["rekening"][0]["catatan"] == expected, (name, data["rekening"][0]["catatan"])
            assert data["kategori"][0]["catatan"] == kategori, (name, data["kategori"][0]["catatan"])

        # Rekening di daftar aset tanpa rincian transaksi tetap muncul, ditandai.
        data = variant("tanpa-rincian", daftar_giro=daftar_lebih)
        assert [(a["norek"], a["catatan"]) for a in data["rekening"]][2] == (
            "5555666677", ["Rincian transaksi rekening ini tidak ditemukan di PDF."]), data["rekening"]
        assert data["kategori"][0]["catatan"] == [] and data["kategori"][0]["rekening"] == 3

        # Judul tabel yang kolomnya tidak dikenal: tidak ada nominal yang ditebak.
        judul_tukar = [CIMB_JUDUL_TABEL[0], ("Transaksi/", "Valuta/", "Uraian Transaksi/", "No Cek/BG", "Kredit/",
                                             "Debit/", "Saldo/"), CIMB_JUDUL_TABEL[2]]
        data = variant("judul", judul_tabel=judul_tukar)
        catatan = data["rekening"][0]["catatan"]
        assert catatan[:2] == ["Hal. 2: susunan kolom tabel mutasi tidak dikenal.",
                               "Hal. 2: teks di luar tabel mutasi — SALDO AWAL 1,000.00"], catatan
        assert "Baris SALDO AWAL tidak ditemukan." in catatan and all(a["catatan"] for a in data["rekening"])
        assert all(not a["transaksi"] for a in data["rekening"])

        # PDF lain dilewati; laporan CIMB tanpa tanggal laporan terbaca gagal dengan pesan jelas.
        write_all(temp / "lain")
        for name in ("bppu", "bpbs", "faktur", "rekening"):
            assert read_cimb(temp / "lain" / f"{name}.pdf") is None, name
        sampul = SimpleNamespace(extract_text=lambda: "CIMB\nCOMBINE STATEMENT PORTFOLIO\nTanggal Laporan : 31 Foo 2025")
        try:
            ExpCore._extract_rekening_cimb([sampul])
            raise AssertionError("tanggal laporan tidak valid harus gagal")
        except ValueError as error:
            assert str(error) == "Tanggal laporan di halaman pertama tidak terbaca.", error

        # Rekap folder: Ringkasan per rekening, Total Kategori per PDF, sheet per rekening; PDF lain dilewati.
        # Januari unik, Februari ada di dua PDF (keduanya ditandai PERLU CEK), Maret berisi selisih.
        folder = temp / "CIMB Ünïcode"
        write_cimb(folder / "a" / "jan.pdf")
        write_cimb(folder / "b" / "feb.pdf", tanggal="28 February 2025")
        write_cimb(folder / "b" / "feb salinan.pdf", tanggal="28 February 2025")
        write_cimb(folder / "selisih.pdf", mutasi=tertukar, tanggal="31 March 2025")
        write_all(folder / "lain")
        (folder / "lain" / "cimb.pdf").unlink()
        (folder / "rusak.pdf").write_bytes(b"corrupt")
        app = ExpCore(log=Mock(), progress=Mock())
        output, summary = app.process_rekening_cimb(str(folder))
        assert output == str(folder / "!Hasil_Rekap_Rekening_CIMB.xlsx")
        assert summary == "Selesai — 4 laporan, 16 rekening, 20 transaksi, 13 perlu dicek, 5 PDF dilewati.", summary
        logs = [call.args[0] for call in app.log.call_args_list]
        assert ("PERINGATAN: laporan CIF 11223344556677 tanggal 28-02-2025 muncul di 2 PDF: feb salinan.pdf, "
                "feb.pdf — ditandai PERLU CEK.") in logs, logs
        assert "DILEWATI: lain\\rekening.pdf — bukan Combine Statement Portfolio CIMB Niaga." in logs, logs
        assert "SESUAI: Total Tabungan Berjangka JANUARI 2025 — 1,680,000.00 (Ekuivalen IDR)" in logs, logs
        assert "  • 05/01: saldo PDF 1,234.50, hasil hitung 1,765.50." in logs, logs

        book = load_workbook(output)
        try:
            assert book.sheetnames[:8] == ["Ringkasan", "Total Kategori", "Semua Transaksi", "1111222233 2025-01",
                                           "4444555566 2025-01", "7777888899 2025-01", "7777000011 2025-01",
                                           "1111222233 2025-02"]
            assert book.sheetnames[11:13] == ["1111222233 2025-02 (2)", "4444555566 2025-02 (2)"], book.sheetnames
            assert len(book.sheetnames) == 19, book.sheetnames
            overview = [[cell.value for cell in row] for row in book["Ringkasan"].iter_rows()]
            assert overview[0] == ["No", "Sheet", "Kategori", "No. Rekening", "Nama Produk", "Mata Uang",
                                   "Tanggal Laporan", "Saldo Awal", "Total Debit", "Total Kredit", "Saldo Akhir",
                                   "Saldo Ekuivalen IDR", "Transaksi Terbaca", "Hasil Cek", "Catatan", "Nasabah",
                                   "Folder Sumber", "File Name"], overview[0]
            assert overview[1] == [1, "1111222233 2025-01", "Tabungan/Giro", "1111222233", "TAPLUS CONTOH", "IDR",
                                   datetime(2025, 1, 31), 1000, 265.5, 500, 1234.5, 1234.5, 2, "SESUAI", None,
                                   "BUDI CONTOH", "a", "jan.pdf"], overview[1]
            assert overview[4][2:4] == ["Tabungan Berjangka", "7777000011"] and overview[4][10:12] == [0, 0]
            assert [row[13] for row in overview[1:]] == ["SESUAI"] * 4 + ["PERLU CEK"] * 9 + ["SESUAI"] * 3
            assert overview[5][6] == datetime(2025, 2, 28) and overview[5][14] == (
                "Laporan periode yang sama juga ada di b\\feb.pdf; transaksinya bisa tercatat ganda."), overview[5]
            assert overview[13][14].startswith("Debit terbaca 0.00 dari 0 transaksi"), overview[13]
            ringkasan = book["Ringkasan"]
            assert ringkasan["B2"].hyperlink.location == "'1111222233 2025-01'!A1"
            assert (ringkasan["D2"].number_format, ringkasan["G2"].number_format, ringkasan["H2"].number_format) == (
                "@", "dd/mm/yyyy", "#,##0.00")
            assert ringkasan.freeze_panes == "A2" and ringkasan.auto_filter.ref == "A1:R17"

            totals = [[cell.value for cell in row] for row in book["Total Kategori"].iter_rows()]
            assert totals[0] == ["No", "Tanggal Laporan", "Kategori", "Jumlah Rekening", "Total PDF (Ekuivalen IDR)",
                                 "Jumlah Saldo Rekening (Ekuivalen IDR)", "Hasil Cek", "Catatan", "Nasabah",
                                 "Folder Sumber", "File Name"], totals[0]
            assert totals[1:3] == [
                [1, datetime(2025, 1, 31), "Tabungan/Giro", 2, 161234.5, 161234.5, "SESUAI", None, "BUDI CONTOH",
                 "a", "jan.pdf"],
                [2, datetime(2025, 1, 31), "Tabungan Berjangka", 2, 1680000, 1680000, "SESUAI", None, "BUDI CONTOH",
                 "a", "jan.pdf"]], totals[1:3]
            assert [row[6] for row in totals[1:]] == ["SESUAI"] * 2 + ["PERLU CEK"] * 4 + ["SESUAI"] * 2
            assert len(totals) == 9

            # Semua Transaksi: seluruh rekening dalam satu tabel, urut kategori lalu nomor rekening, dan tetap
            # kronologis per rekening (laporan diurutkan per tanggal); kolom Sheet menautkan ke sheet asalnya.
            semua = book["Semua Transaksi"]
            rows = [[cell.value for cell in row] for row in semua.iter_rows()]
            assert rows[0] == ["No", "Tanggal Transaksi", "Tanggal Valuta", "Kategori", "No. Rekening", "Nama Produk",
                               "Mata Uang", "Uraian Transaksi", "No Cek/BG", "Debit", "Kredit", "Saldo", "Hasil Cek",
                               "Tanggal Laporan", "Sheet", "File Name"], rows[0]
            assert len(rows) == 21, len(rows)
            sheets = lambda norek, count: [(norek, f"{norek} {suffix}")
                                           for suffix in ("2025-01", "2025-02", "2025-02 (2)", "2025-03")
                                           for _ in range(count)]
            assert [(r[4], r[14]) for r in rows[1:]] == (
                sheets("1111222233", 2) + sheets("7777000011", 1) + sheets("7777888899", 2)), rows
            assert rows[1] == [1, datetime(2024, 12, 31), datetime(2024, 12, 31), "Tabungan/Giro", "1111222233",
                               "TAPLUS CONTOH", "IDR", "SETORAN TUNAI CABANG CONTOH REF0001", None, None, 500, 1500,
                               "SESUAI", datetime(2025, 1, 31), "1111222233 2025-01", "jan.pdf"], rows[1]
            assert rows[2][8:12] == ["CK001", 265.5, None, 1234.5], rows[2]
            # Hasil Cek mengikuti rekeningnya: Februari ganda dan rekening Maret yang selisih ikut PERLU CEK.
            assert [r[12] for r in rows[1:]] == (["SESUAI"] * 2 + ["PERLU CEK"] * 6 + ["SESUAI"] + ["PERLU CEK"] * 2
                                                 + ["SESUAI"] * 3 + ["PERLU CEK"] * 4 + ["SESUAI"] * 2), rows
            assert rows[13][3:7] == ["Tabungan Berjangka", "7777888899", "RENCANA CONTOH USD", "USD"], rows[13]
            assert semua["O2"].hyperlink.location == "'1111222233 2025-01'!A1"
            assert (semua["B2"].number_format, semua["E2"].number_format, semua["I3"].number_format,
                    semua["J3"].number_format, semua["N2"].number_format) == (
                "dd/mm/yyyy", "@", "@", "#,##0.00", "dd/mm/yyyy")
            assert semua.freeze_panes == "A2" and semua.auto_filter.ref == "A1:P21"

            sheet = book["1111222233 2025-01"]
            table = [[cell.value for cell in row[:7]] for row in sheet.iter_rows(max_row=3)]
            assert table == [
                ["TANGGAL TRANSAKSI", "TANGGAL VALUTA", "URAIAN TRANSAKSI", "NO CEK/BG", "DEBIT", "KREDIT", "SALDO"],
                [datetime(2024, 12, 31), datetime(2024, 12, 31), "SETORAN TUNAI CABANG CONTOH REF0001", None, None,
                 500, 1500],
                [datetime(2025, 1, 5), datetime(2025, 1, 6), "TARIK TUNAI", "CK001", 265.5, None, 1234.5]], table
            assert sheet["A4"].value is None, "ringkasan tidak boleh di bawah tabel"
            assert (sheet["A2"].number_format, sheet["B2"].number_format, sheet["D3"].number_format,
                    sheet["E3"].number_format) == ("dd/mm/yyyy", "dd/mm/yyyy", "@", "#,##0.00")
            assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:G3"
            panel = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=1, max_row=19, min_col=9, max_col=10)]
            assert panel == [
                ["INFORMASI REKENING", None], ["Kategori", "Tabungan/Giro"], ["No. Rekening", "1111222233"],
                ["Nama Produk", "TAPLUS CONTOH"], ["Mata Uang", "IDR"], ["Nasabah", "BUDI CONTOH"],
                ["Tanggal Laporan", datetime(2025, 1, 31)], ["File", "jan.pdf"], [None, None],
                ["RINGKASAN PDF", "NILAI"], ["SALDO AWAL", 1000], ["TOTAL DEBIT", 265.5], ["TOTAL KREDIT", 500],
                ["SALDO AKHIR", 1234.5], ["SALDO DI DAFTAR ASET", 1234.5], ["SALDO EKUIVALEN IDR", 1234.5],
                [None, None], ["HASIL CEK", None], ["SESUAI", None]], panel
            assert sheet["J3"].number_format == "@" and sheet["J11"].number_format == "#,##0.00"
            flagged = book["1111222233 2025-03"]
            assert flagged["I19"].value == "PERLU CEK" and flagged["I20"].value.startswith("• Debit terbaca")
            ganda = book["1111222233 2025-02"]
            assert (ganda["I19"].value, ganda["I20"].value) == (
                "PERLU CEK", "• Laporan periode yang sama juga ada di b\\feb.pdf; transaksinya bisa tercatat ganda.")
        finally:
            book.close()

        empty = temp / "bukan-cimb"
        write_all(empty)
        (empty / "cimb.pdf").unlink()
        output, summary = ExpCore(log=Mock(), progress=Mock()).process_rekening_cimb(str(empty))
        assert output is None and summary.startswith("Tidak ada laporan CIMB Niaga yang cocok."), summary
    print("Rekening koran CIMB: sel multi-baris, kolom dari garis tabel, lintas halaman, cek total & saldo, Excel: ok")


def test_pdf_samples(perbarui_acuan=False):
    """Regresi opsional pada PDF contoh lokal di contoh-pdf/ (PDF asli tidak diubah).

    Excel setiap pekerjaan dibandingkan per sel dengan acuan di contoh-pdf/acuan; perbarui_acuan menimpanya.
    """
    import shutil

    root = Path(__file__).parent / "contoh-pdf"

    def text(path):
        with pdfplumber.open(path) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)

    # Satu PDF per folder; nama berkasnya bisa berubah (mis. sudah diberi nama oleh ExpCore).
    bpbs, = sorted((root / "BUPOT 2024").glob("*.pdf"))
    bppu, = sorted((root / "BUPOT 2026").glob("*.pdf"))
    assert len(ExpCore._extract_bupot2024_rows(text(bpbs))) == 1
    assert len(ExpCore._extract_bupot_rows(text(bppu))) == 1
    # Penamaan Bupot membaca keduanya dalam satu folder: BPPU (C.3/A.2) dan BPBS (C.2/A.3).
    expected = {
        "pemotong": {"BPPU": "ADIRA DINAMIKA MULTI FINANCE TBK - 2507UMKHM - 12-2025 - TIDAK FINAL - NORMAL.pdf",
                     "BPBS": "BFI FINANCE INDONESIA TBK - 2000000001 - 01-2024 - TIDAK FINAL - NORMAL.pdf"},
        "penerima": {"BPPU": "MITRACOLL SARANA JAYA - 2507UMKHM - 12-2025 - TIDAK FINAL - NORMAL.pdf",
                     "BPBS": "MITRACOLL SARANA JAYA - 2000000001 - 01-2024 - TIDAK FINAL - NORMAL.pdf"},
    }
    with tempfile.TemporaryDirectory() as folder:
        shutil.copy(bppu, folder)
        shutil.copy(bpbs, folder)
        for source, names in expected.items():
            log_path, _ = ExpCore().process_rename_bupot(folder, name_source=source)
            with open(log_path, encoding="utf-8-sig", newline="") as log:
                rows = {row["FORMULIR"]: row for row in csv.DictReader(log)}
            assert {form: row["nama_baru"] for form, row in rows.items()} == names, rows
            assert all(row["status"] in ("SIAP", "SUDAH SESUAI") for row in rows.values()), rows
    print("PDF contoh Bupot 2024/2026: rekap dan penamaan BPPU + BPBS (C.3/C.2 dan A.2/A.3): ok", flush=True)

    # Jumlah transaksi per bulan; semua bulan harus SESUAI dan saldo akhir = saldo awal bulan berikutnya.
    for folder, counts in (("Rekening Giro", [28, 29, 20]), ("Rekening Tahapan", [81, 105, 101])):
        statements = [read_rekening(path) for path in sorted((root / folder).glob("*.pdf"))]
        assert [len(s["transaksi"]) for s in statements] == counts, (folder, [len(s["transaksi"]) for s in statements])
        assert all(not s["catatan"] for s in statements), [s["catatan"] for s in statements]
        for before, after in zip(statements, statements[1:]):
            assert before["ringkasan"]["SALDO AKHIR"][0] == after["ringkasan"]["SALDO AWAL"][0], folder
        print(f"PDF contoh {folder}: {len(statements)} bulan, {sum(counts)} transaksi, semua SESUAI", flush=True)

    # CIMB Niaga Jul–Des 2024: transaksi per laporan, semua rekening & total kategori SESUAI, dan saldo akhir
    # tiap rekening = saldo awal bulan berikutnya (bila rekening ada di kedua bulan).
    reports = [read_cimb(path) for path in sorted((root / "BANG CIMB").glob("*.pdf"))]
    counts = [sum(len(a["transaksi"]) for a in r["rekening"]) for r in reports]
    assert counts == [29, 46, 36, 32, 41, 24], counts
    assert [(r["bulan"], len(r["rekening"])) for r in reports] == [(7, 4), (8, 5), (9, 3), (10, 3), (11, 4), (12, 3)]
    assert all(not item["catatan"] for r in reports for item in r["rekening"] + r["kategori"]), \
        [(r["periode"], item["catatan"]) for r in reports for item in r["rekening"] + r["kategori"] if item["catatan"]]
    for before, after in zip(reports, reports[1:]):
        akhir = {a["norek"]: a["SALDO AKHIR"] for a in before["rekening"]}
        for akun in after["rekening"]:
            assert akun["norek"] not in akhir or akhir[akun["norek"]] == akun["SALDO AWAL"], (after["periode"], akun["norek"])
    print(f"PDF contoh CIMB Niaga: {len(reports)} laporan, {sum(counts)} transaksi, semua SESUAI", flush=True)

    for nama, (metode, keluaran) in SAMPLE_EXPORTS.items():
        sumber, acuan = root / nama, ACUAN / nama / keluaran
        with tempfile.TemporaryDirectory() as folder:
            for pdf in sorted(sumber.glob("*.pdf")):
                shutil.copy(pdf, folder)
            hasil, _ = getattr(ExpCore(), metode)(folder)
            assert hasil, (nama, "tidak menghasilkan Excel")
            if perbarui_acuan:
                acuan.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy(hasil, acuan)
                print(f"Acuan Excel {nama}: diperbarui", flush=True)
                continue
            assert acuan.is_file(), (f"Acuan {acuan} belum ada. Buat dengan: "
                                     "./.venv/Scripts/python.exe test_expcore.py --pdf-samples --perbarui-acuan")
            beda = bandingkan_excel(acuan, hasil)
            assert not beda, f"{nama}: Excel berbeda dari acuan\n" + "\n".join(beda)
        print(f"Acuan Excel {nama}: identik per sel", flush=True)


def main():
    test_batch_exports()
    test_rename_pemotong()
    test_rename_bpbs()
    test_rekening_koran()
    test_rekening_cimb()
    test_engine_protocol()
    text = """2505Z0UR6 10-2025 TIDAK FINAL PEMBETULAN KE-2
C.3 NAMA PEMOTONG : PT CONTOH: ABADI
C.4 TANGGAL : 7 Juli 2026"""
    assert ExpCore._extract_rename_bupot_data(text) == {
        "NAMA_PENERIMA": "",
        "NAMA_PEMOTONG": "PT CONTOH: ABADI",
        "NOMOR_BUKTI": "2505Z0UR6",
        "MASA_PAJAK": "10-2025",
        "SIFAT": "TIDAK FINAL",
        "STATUS": "PEMBETULAN KE-2",
        "FORMULIR": "",
    }
    assert ExpCore._safe_filename("PT: CONTOH?") == "PT CONTOH"

    fallback = """ABC12345
11-2025
FINAL
DIBATALKAN
C.3 NAMA PEMOTONG : CV FALLBACK
C.4 TANGGAL : 7 Juli 2026"""
    assert ExpCore._extract_rename_bupot_data(fallback) == {
        "NAMA_PENERIMA": "",
        "NAMA_PEMOTONG": "CV FALLBACK",
        "NOMOR_BUKTI": "ABC12345",
        "MASA_PAJAK": "11-2025",
        "SIFAT": "FINAL",
        "STATUS": "DIBATALKAN",
        "FORMULIR": "",
    }

    with tempfile.TemporaryDirectory() as folder:
        path = os.path.join(folder, "Bupot.pdf")
        open(path, "w").close()
        assert ExpCore._unique_filename(path).endswith("Bupot (2).pdf")

    print("Penamaan Otomatis Bupot: ok")

    # Label "PEMUNGUT PPh" terpotong baris — sisa "PPh" tidak boleh ikut nama.
    wrapped = """25004WOBY 01-2025 TIDAK FINAL NORMAL
C.3 NAMA PEMOTONG DAN/ATAU PEMUNGUT : ADIRA DINAMIKA MULTI FINANCE TBK.
PPh
C.4 TANGGAL : 31 Januari 2025"""
    assert ExpCore._extract_rename_bupot_data(wrapped)["NAMA_PEMOTONG"] == \
        "ADIRA DINAMIKA MULTI FINANCE TBK."

    bupot = """25004WOBY 01-2025 TIDAK FINAL NORMAL
A.1 NPWP / NIK : 0436117337418000
A.2 NAMA : MITRACOLL SARANA JAYA
A.3 NOMOR IDENTITAS : 0436117337418000000000
B.1 Jenis Fasilitas : Tanpa Fasilitas
B.2 Jenis PPh : Pasal 23
B.3 B.4 B.5 B.6 B.7
24-104-18 Jasa Perantara dan/atau Keagenan 3.700.000 2 74.000
B.8 Dokumen Dasar Bukti Jenis Dokumen : Surat Tagihan Tanggal : 31 Januari 2025
Pemotongan dan/atau Pemungutan PPh Unifikasi
B.9 Nomor Dokumen : 10006302
B.10 Untuk Instansi Pemerintah
C.1 NPWP / NIK : 0013464946091000
C.3 NAMA PEMOTONG DAN/ATAU PEMUNGUT : ADIRA DINAMIKA MULTI FINANCE TBK.
PPh
C.4 TANGGAL : 31 Januari 2025
C.5 NAMA PENANDATANGAN : I DEWA MADE SUSILA"""
    rows = ExpCore._extract_bupot_rows(bupot)
    assert len(rows) == 1, rows
    assert rows[0] == {
        "Nomor Dokumen": "25004WOBY", "Masa Pajak": "01-2025",
        "NPWP/NIK": "0436117337418000", "Nama": "MITRACOLL SARANA JAYA",
        "Status Bukti": "NORMAL", "Jenis Fasilitas": "Tanpa Fasilitas",
        "Jenis PPh": "Pasal 23", "Kode Objek Pajak": "24-104-18",
        "Objek Pajak": "Jasa Perantara dan/atau Keagenan",
        "DPP (Rp)": 3700000.0, "Tarif (%)": 2.0, "Pajak Penghasilan (Rp)": 74000.0,
        "Jenis Dokumen": "Surat Tagihan", "Nomor Dokumen Dasar": "10006302",
        "Tanggal Dokumen": "31 Januari 2025",
        "NPWP/NIK Pemotong": "0013464946091000",
        "Nama Pemotong": "ADIRA DINAMIKA MULTI FINANCE TBK.",
        "Tanggal Bukti Potong": "31 Januari 2025",
    }
    # Baris objek pajak tanpa nominal: PDF gagal dengan pesan, bukan baris DPP/PPh 0.
    try:
        ExpCore._extract_bupot_rows(bupot.replace(" 3.700.000 2 74.000", ""))
    except ValueError as error:
        assert "tidak terbaca" in str(error), error
    else:
        raise AssertionError("nominal yang tidak terbaca harus menggagalkan PDF, bukan menjadi 0")

    print("Rekap Bukti Potong: ok")

    bupot2024 = """a r e a s t a p l e s
BUKTI PEMOTONGAN/PEMUNGUTAN
FORMULIR BPBS
PPh PASAL 4 AYAT (2), PASAL 15, PASAL 22, DAN PASAL 23
H.1 NOMOR : 2 0 0 0 0 0 0 0 0 1 H.4 PPh Final
KEMENTERIAN KEUANGAN RI
DIREKTORAT JENDERAL PAJAK H.2 Pembetulan Ke- 0 H.3 Pembatalan H.5 X PPh Tidak Final
A. IDENTITAS WAJIB PAJAK YANG DIPOTONG/DIPUNGUT
A.1 NPWP : 4 3 6 1 1 7 3 3 7 4 1 8 0 0 0
A.2 NIK :
A.3 Nama : MITRACOLL SARANA JAYA
B. PAJAK PENGHASILAN YANG DIPOTONG/DIPUNGUT
Masa Pajak Dikenakan Tarif Lebih PPh yang Dipotong/
(mm-yyyy) Kode Objek Pajak Dasar Pengenaan Pajak (Rp) m Ti e n m g i g li i k ( i T N id P a W k P) Tarif (%) Dipungut/DTP (Rp)
B.1 B.2 B.3 B.4 B.5 B.6
1-2024 24-104-18 1.300.000,00 2.00 26.000,00
Keterangan Kode Objek Pajak : Jasa Perantara dan/atau Keagenan
B.7 Dokumen Referensi : Nomor Dokumen 1.3-01.24-0000407
Nama Dokumen Surat Perjanjian Tanggal 3 1 dd 0 1 mm 2 0 2 4 yyyy
B.8 Dokumen Referensi untuk Faktur Pajak, apabila ada :
Nomor Faktur Pajak : Tanggal dd mm yyyy
B.9 PPh dibebankan berdasarkan Surat Keterangan Bebas (SKB).
Nomor : Tanggal dd mm yyyy
B.10 PPh yang ditanggung oleh Pemerintah (DTP) berdasarkan :
B.11 PPh dalam hal transaksi menggunakan Surat Keterangan berdasarkan PP Nomor 23 Tahun 2018 dengan Nomor :
B.12 PPh yang dipotong/dipungut yang diberikan fasilitas PPh berdasarkan:
C. IDENTITAS PEMOTONG/PEMUNGUT
C.1 NPWP : 0 1 3 1 6 1 1 1 2 5 0 1 0 0 1
C.2 Nama Wajib Pajak : BFI FINANCE INDONESIA TBK.
C.3 Tanggal : 3 1 dd 0 1 mm 2 0 2 4 yyyy
C.4 Nama Penandatangan : ARIEF LUTHFI BACHTIAR
C.5 Pernyataan Wajib Pajak : Dengan ini saya menyatakan bahwa bukti Pemotongan/Pemungutan Unifikasi telah saya isi dengan benar dan telah saya tandatangani secara
elektronik
Apabila terdapat kesalahan/pembatalan dalam pembuatan Bukti Pemotongan/Pemungutan Unifikasi yang menyebabkan kelebihan
pemotongan/pemungutan PPh atau pembayaran, atas kelebihan tersebut akan diajukan:
Pengembalian atas kelebihan pembayaran pajak yang tidak seharusnya terutang oleh Pemotong dan/atau Pemungut PPh
V Pemindahbukuan oleh Pemotong dan/atau Pemungut PPh
Sesuai dengan ketentuan yang berlaku di, Direktorat Jenderal pajak mengatur bahwa Bukti Pemotongan/Pemungutan PPh Unifikasi ini
9PT5ZJ6S dinyatakan sah dan tidak diperlukan tanda tangan basah pada Bukti Pemotongan ini."""
    rows2024 = ExpCore._extract_bupot2024_rows(bupot2024)
    assert len(rows2024) == 1, rows2024
    assert rows2024[0] == {
        "Nomor Bukti Potong": "2000000001",
        "Pembetulan Ke": 0,
        "Status Bukti": "NORMAL",
        "Sifat": "TIDAK FINAL",
        "NPWP": "436117337418000",
        "NIK": "-",
        "Nama": "MITRACOLL SARANA JAYA",
        "Masa Pajak": "01-2024",
        "Kode Objek Pajak": "24-104-18",
        "Objek Pajak": "Jasa Perantara dan/atau Keagenan",
        "DPP (Rp)": 1300000.0,
        "Tarif Lebih Tinggi": "Tidak",
        "Tarif (%)": 2.0,
        "Pajak Penghasilan (Rp)": 26000.0,
        "Nomor Dokumen Referensi": "1.3-01.24-0000407",
        "Nama Dokumen": "Surat Perjanjian",
        "Tanggal Dokumen": "31-01-2024",
        "NPWP Pemotong": "013161112501001",
        "Nama Pemotong": "BFI FINANCE INDONESIA TBK.",
        "Tanggal Bukti Potong": "31-01-2024",
        "Nama Penandatangan": "ARIEF LUTHFI BACHTIAR",
    }

    # Tarif BPBS bisa bulat, desimal titik, atau desimal koma; centang X opsional.
    for tarif, expected_tarif in (("2", 2.0), ("2.00", 2.0), ("2,00", 2.0),
                                  ("1.5", 1.5), ("1,5", 1.5), ("0", 0.0), ("10", 10.0)):
        for checkbox in ("", "X "):
            variant = bupot2024.replace("2.00 26.000,00", f"{checkbox}{tarif} 26.000,00")
            expected_row = dict(rows2024[0])
            expected_row["Tarif (%)"] = expected_tarif
            expected_row["Tarif Lebih Tinggi"] = "Ya" if checkbox else "Tidak"
            assert ExpCore._extract_bupot2024_rows(variant) == [expected_row], (tarif, checkbox)

    # Tarif kosong/rusak tidak boleh membuat nominal PPh terbaca sebagai tarif.
    for invalid_tarif in ("", "2.", "2,", "abc"):
        variant = bupot2024.replace("2.00 26.000,00", f"{invalid_tarif} 26.000,00")
        assert ExpCore._extract_bupot2024_rows(variant) == [], invalid_tarif

    # Formulir Coretax (BPPU) tidak boleh ikut terbaca oleh parser BPBS.
    assert ExpCore._extract_bupot2024_rows(bupot) == []
    # Tanggal dengan kotak kosong menghasilkan string kosong, bukan crash.
    assert ExpCore._form_date("Tanggal dd mm yyyy") == ""
    assert ExpCore._form_date("3 1 dd 0 1 mm 2 0 2 4 yyyy") == "31-01-2024"

    print("Rekap Bukti Potong 2024: ok")


if __name__ == "__main__":
    main()
    if "--pdf-samples" in sys.argv:
        test_pdf_samples(perbarui_acuan="--perbarui-acuan" in sys.argv)
